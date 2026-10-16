from docx import Document

doc = Document('weekly_report_source.docx')
for i, table in enumerate(doc.tables):
    print(f'\nTABLE {i} ({len(table.rows)}x{len(table.columns)})')
    for j, row in enumerate(table.rows):
        print(j, ' | '.join(cell.text.replace('\n', ' / ') for cell in row.cells))

print('\nPARAGRAPHS')
for i, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.strip():
        print(i, repr(paragraph.text))
